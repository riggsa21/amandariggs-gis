"""GSI sample automation. Synthetic data only, no client data.
Read a recurring order export, normalize dates/currency/status, flag exceptions,
and generate Excel workbooks without changing the original input file.
Requires: pip install openpyxl
"""
from pathlib import Path
import csv
from datetime import datetime
from decimal import Decimal, InvalidOperation
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parent
INPUT=ROOT/'incoming_orders.csv'
OUT=ROOT/'output'; OUT.mkdir(exist_ok=True)

def parse_date(raw):
    for fmt in ('%m/%d/%Y','%Y-%m-%d','%m-%d-%Y','%b %d, %Y'):
        try:return datetime.strptime(raw.strip(),fmt).date()
        except ValueError:pass
    raise ValueError('unrecognized date')

def make_book(filename,header,rows,amount_column=None):
    wb=Workbook();ws=wb.active;ws.title='Results'
    ws.append(header)
    for row in rows: ws.append(row)
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid', fgColor='174C59')
        cell.alignment=Alignment(horizontal='left')
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for ci in range(1,len(header)+1):
        values=[str(ws.cell(ri,ci).value or '') for ri in range(1,ws.max_row+1)]
        ws.column_dimensions[get_column_letter(ci)].width=min(36,max(15,max(map(len,values))+3))
    if amount_column:
        for ri in range(2,ws.max_row+1): ws.cell(ri,amount_column).number_format='$#,##0.00'
    wb.save(OUT/filename)

with INPUT.open(newline='',encoding='utf-8-sig') as fh:
    original=list(csv.DictReader(fh))
clean=[];review=[]; seen=set()
for source in original:
    oid=source['Order ID'].strip()
    try:
        if not oid: raise ValueError('missing order id')
        if oid in seen: raise ValueError('duplicate order id')
        dt=parse_date(source['Order Date'])
        amount=Decimal(source['Amount'].replace('$','').replace(',','').strip())
        if not amount.is_finite() or amount < 0:raise ValueError('invalid amount')
        customer=source['Customer'].strip()
        if not customer:raise ValueError('missing customer')
        status=source['Status'].strip().lower()
        clean.append((oid,dt.isoformat(),customer,float(amount),status))
        seen.add(oid)
    except (ValueError,InvalidOperation) as ex:
        review.append((oid,source['Order Date'],source['Customer'],source['Amount'],source['Status'],str(ex)))

make_book('01_Clean_Orders.xlsx',['Order ID','Order Date','Customer','Amount','Status'],clean,4)
make_book('02_Review_Needed.xlsx',['Order ID','Order Date','Customer','Amount','Status','Reason'],review)
by_customer=Counter()
for _,_,name,amount,_ in clean:by_customer[name]+=round(amount*100)
summary=[(name,round(cents/100,2)) for name,cents in sorted(by_customer.items())]
make_book('03_Customer_Summary.xlsx',['Customer','Total Amount'],summary,2)
report=f'''GSI WORKFLOW DEMO | FICTIONAL SAMPLE DATA
Input rows:       {len(original)}
Clean orders:     {len(clean)}
Flagged for QA:   {len(review)}
Summary groups:   {len(summary)}
Source preserved: Yes
'''
(OUT/'Run_Summary.txt').write_text(report,encoding='utf-8')
print('\n'+'='*46);print(' GEOSITE INTEL | ORDER PROCESSING AUTOMATION');print('='*46)
print(report)
for p in sorted(OUT.iterdir()):print('Created:',p.name)
print('\nDONE: Repeatable, documented output. Human review still required for flags.')
